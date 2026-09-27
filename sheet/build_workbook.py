"""
Builds ED_PAC_Register.xlsx — the register in the exact layout Code.gs reads,
with dropdowns, live duration formulas, a Reference sheet holding the bed
establishment, and the full-capacity demonstration scenario (217 records).

Mirrors SetupSheet.gs. Regenerate with:  python3 build_workbook.py
"""
import random, datetime
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

random.seed(20250305)   # reproducible output

HEADERS = ['Location','Triage Date/Time','Full Name','Initial','IC / Passport','MRN',
           'Age','Gender','Zone Code','Bed / Position Code','Current Zone','Status',
           'Referred To','Queue No. (GZ only)','Called into GZ Room',
           'Pre-admit Date/Time','Admit Date/Time','BWT (hh:mm)','TWT (hh:mm)',
           'GZWT (hh:mm)','Discharge Date/Time']

LOCATIONS   = ['ED WCC','ED BU','PAC WCC']
GENDERS     = ['Male','Female']
ZONE_CODES  = ['rz','yz','gz','ob','ab','pac']
STATUSES    = ['ongoingtreatment','referred','preadmit','admitted','discharge']
DISCIPLINES = ['Medical','Surgical','Orthopedic','O&G','Paediatric',
               'Paediatric Dental','OMFS','Psychiatry','Special Needs Dental']
ZONE_NAMES  = {'rz':'Red Zone','yz':'Yellow Zone','gz':'Green Zone',
               'ob':'Observation Bay','ab':'Asthma Bay','pac':'Patient Assessment Centre'}

# location, zone, prefix, funded, crisisTo, unit, waiting
ESTABLISHMENT = [
    ('ED WCC','rz','wccrz',   4, 10, 'bed', 0),
    ('ED WCC','yz','wccyz',   4, 12, 'bed', 0),
    ('ED WCC','ob','wccob',   8, 10, 'bed', 0),
    ('ED WCC','ab','wccab',   4,  4, 'sofa', 0),
    ('ED WCC','gz','wccgz',   2,  2, 'consultation room', 50),
    ('ED BU','rz','burz',     6, 12, 'bed', 0),
    ('ED BU','yz','buyz',    16, 50, 'bed', 0),
    ('ED BU','gz','bugz',     2,  2, 'consultation room', 50),
    ('PAC WCC','pac','wccpac',8, 15, 'bed', 0),
]

TEAL, TEAL_LT, AMBER_LT = '055257', 'E0F5F6', 'FFF4E0'

def p2(n): return f'{n:02d}'

def inventory():
    out = []
    for loc, zone, prefix, funded, crisis_to, unit, waiting in ESTABLISHMENT:
        for n in range(1, funded + 1):
            out.append(dict(code=f'{prefix}{p2(n)}', location=loc, zone=zone, kind='funded',
                            unit=unit, desc=f'{loc} {ZONE_NAMES[zone]} {unit} {p2(n)}'))
        for c in range(funded + 1, crisis_to + 1):
            out.append(dict(code=f'{prefix}{p2(c)}crisis', location=loc, zone=zone, kind='crisis',
                            unit=unit,
                            desc=f'{loc} {ZONE_NAMES[zone]} {unit} {p2(c)} — escalation capacity, crisis mode active'))
        if waiting:
            out.append(dict(code=f'{prefix}-waiting', location=loc, zone=zone, kind='waiting',
                            unit='queue position',
                            desc=f'{loc} {ZONE_NAMES[zone]} waiting area — holds a queue number, not a bed'))
    return out

FIRST = ['Ahmad','Aisyah','Amirah','Anis','Azman','Danial','Faizal','Farah','Hafiz',
         'Hafizah','Hamizan','Hasnah','Helmi','Izzuddin','Mohamed','Nizam','Norizan',
         'Norliza','Nurul','Rizal','Rohani','Shahril','Siti','Suraya','Syed','Zainab',
         'Zulkifli','Zuraidah']
LAST = ['Abdullah','Ahmad','Hamid','Hassan','Ibrahim','Ismail','Mohamed','Omar',
        'Osman','Rahman','Salleh','Yusof','Zakaria']
IC_STATES = [f'{i:02d}' for i in range(1, 15)]

CLOCK = datetime.datetime(2025, 3, 5, 15, 0)
OPEN_FROM = CLOCK - datetime.timedelta(hours=8)

def build_rows():
    slots = []
    for b in inventory():
        if b['kind'] == 'waiting':
            waiting = next(e[6] for e in ESTABLISHMENT
                           if e[2] == b['code'].replace('-waiting', ''))
            for q in range(1, waiting + 1):
                slots.append((b, q))
        else:
            slots.append((b, None))

    used_ic, rows = set(), []
    span = (CLOCK - OPEN_FROM).total_seconds()
    for bed, queue_no in slots:
        waiting = bed['kind'] == 'waiting'
        room = bed['unit'] == 'consultation room'
        frac = (0.35 + random.random() * 0.65) if waiting else random.random()
        triage = (OPEN_FROM + datetime.timedelta(seconds=int(frac * span))).replace(second=0, microsecond=0)

        first, last = random.choice(FIRST), random.choice(LAST)
        female = True if bed['location'] == 'PAC WCC' else random.random() < 0.52
        name = f"{first} {'binti' if female else 'bin'} {last}"

        if bed['location'] == 'PAC WCC':
            age = random.randint(18, 41)
        elif bed['location'] == 'ED WCC':
            age = random.randint(1, 17) if random.random() < 0.55 else random.randint(18, 67)
        else:
            age = random.randint(16, 85)

        while True:
            yy = (2025 - age) % 100
            ic = f'{yy:02d}{random.randint(1,12):02d}{random.randint(1,28):02d}-' \
                 f'{random.choice(IC_STATES)}-{random.randint(0,9999):04d}'
            if ic not in used_ic:
                used_ic.add(ic); break

        r = {h: '' for h in HEADERS}
        r['Location'] = bed['location']; r['Triage Date/Time'] = triage
        r['Full Name'] = name; r['Initial'] = first
        r['IC / Passport'] = ic; r['MRN'] = f'MRN{random.randint(100000,999999)}'
        r['Age'] = age; r['Gender'] = 'Female' if female else 'Male'
        r['Zone Code'] = bed['zone']; r['Bed / Position Code'] = bed['code']
        r['Current Zone'] = bed['zone']

        if waiting:
            r['Status'] = 'ongoingtreatment'
            r['Queue No. (GZ only)'] = str(queue_no)
        elif room:
            r['Status'] = 'ongoingtreatment'
            r['Queue No. (GZ only)'] = str(random.randint(1, 200))
            called = triage + datetime.timedelta(minutes=random.randint(20, 140))
            if called < CLOCK: r['Called into GZ Room'] = called
        else:
            roll = random.random()
            status = ('ongoingtreatment' if roll < 0.56 else 'referred' if roll < 0.72
                      else 'preadmit' if roll < 0.86 else 'admitted' if roll < 0.96 else 'discharge')
            r['Status'] = status
            if status in ('referred','preadmit','admitted'):
                r['Referred To'] = random.choice(DISCIPLINES)
            if status in ('preadmit','admitted'):
                pre = triage + datetime.timedelta(minutes=random.randint(60, 300))
                if pre > CLOCK: pre = CLOCK - datetime.timedelta(minutes=10)
                r['Pre-admit Date/Time'] = pre
                if status == 'admitted':
                    adm = pre + datetime.timedelta(minutes=random.randint(40, 180))
                    r['Admit Date/Time'] = min(adm, CLOCK)
            if status == 'discharge':
                dis = triage + datetime.timedelta(minutes=random.randint(90, 290))
                r['Discharge Date/Time'] = min(dis, CLOCK)
        rows.append(r)

    rows.sort(key=lambda x: x['Triage Date/Time'])
    return rows


wb = Workbook()
ws = wb.active; ws.title = 'Sheet1'
inv = inventory()
rows = build_rows()
n = len(rows)
last_col = len(HEADERS)
DATE_FMT = 'dd/mm/yyyy hh:mm'

# Row 1 — title band
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
t = ws.cell(1, 1, 'EMERGENCY DEPARTMENT & PAC — PATIENT TRACKING REGISTER   |   '
                  'Hospital Tengku Permaisuri Norashikin, Kajang')
t.font = Font(bold=True, size=12, color='FFFFFF')
t.fill = PatternFill('solid', fgColor=TEAL)
t.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 32

# Row 2 — headers
thin = Side(style='thin', color='B8D8DA')
for c, h in enumerate(HEADERS, start=1):
    cell = ws.cell(2, c, h)
    cell.font = Font(bold=True, color=TEAL)
    cell.fill = PatternFill('solid', fgColor=TEAL_LT)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = Border(bottom=thin, top=thin, left=thin, right=thin)
ws.row_dimensions[2].height = 42
ws.freeze_panes = 'B3'

DATE_COLS = ['Triage Date/Time','Called into GZ Room','Pre-admit Date/Time',
             'Admit Date/Time','Discharge Date/Time']
TEXT_COLS = ['IC / Passport','Queue No. (GZ only)','BWT (hh:mm)','TWT (hh:mm)','GZWT (hh:mm)']

for i, r in enumerate(rows):
    excel_row = 3 + i
    for c, h in enumerate(HEADERS, start=1):
        v = r[h]
        if h in ('BWT (hh:mm)','TWT (hh:mm)','GZWT (hh:mm)'):
            continue                                    # formulas written below
        cell = ws.cell(excel_row, c, v if v != '' else None)
        if h in DATE_COLS and v != '':
            cell.number_format = DATE_FMT
        elif h in TEXT_COLS:
            cell.number_format = '@'
        elif h == 'Age':
            cell.number_format = '0'

    # Live duration formulas, identical to the ones SetupSheet.gs installs.
    ws.cell(excel_row, 18, f'=IF(OR(Q{excel_row}="",P{excel_row}=""),"",'
                           f'TEXT(Q{excel_row}-P{excel_row},"[h]:mm"))')
    ws.cell(excel_row, 19, f'=IF(B{excel_row}="","",IF(Q{excel_row}<>"",'
                           f'TEXT(Q{excel_row}-B{excel_row},"[h]:mm"),'
                           f'IF(U{excel_row}<>"",TEXT(U{excel_row}-B{excel_row},"[h]:mm"),"")))')
    ws.cell(excel_row, 20, f'=IF(OR(O{excel_row}="",B{excel_row}=""),"",'
                           f'TEXT(O{excel_row}-B{excel_row},"[h]:mm"))')
    for c in (18, 19, 20):
        ws.cell(excel_row, c).alignment = Alignment(horizontal='center')

widths = [11,18,24,12,17,13,6,9,10,21,13,17,19,13,18,18,18,11,11,11,18]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

# ── Reference sheet ────────────────────────────────────────
ref = wb.create_sheet('Reference')
for c, h in enumerate(['Position Code','Location','Zone','Zone Name','Kind','Description'], start=1):
    cell = ref.cell(1, c, h)
    cell.font = Font(bold=True, color=TEAL); cell.fill = PatternFill('solid', fgColor=TEAL_LT)
for i, b in enumerate(inv, start=2):
    ref.cell(i, 1, b['code']); ref.cell(i, 2, b['location']); ref.cell(i, 3, b['zone'])
    ref.cell(i, 4, ZONE_NAMES[b['zone']]); ref.cell(i, 5, b['kind']); ref.cell(i, 6, b['desc'])

SUM_COL = 8
for c, h in enumerate(['Location','Zone','Unit','Funded capacity','Escalation beds','Waiting places'],
                      start=SUM_COL):
    cell = ref.cell(1, c, h)
    cell.font = Font(bold=True, color='8A4B00'); cell.fill = PatternFill('solid', fgColor=AMBER_LT)
for i, (loc, zone, prefix, funded, crisis_to, unit, waiting) in enumerate(ESTABLISHMENT, start=2):
    ref.cell(i, SUM_COL, loc); ref.cell(i, SUM_COL+1, ZONE_NAMES[zone])
    ref.cell(i, SUM_COL+2, unit); ref.cell(i, SUM_COL+3, funded)
    ref.cell(i, SUM_COL+4, max(0, crisis_to - funded)); ref.cell(i, SUM_COL+5, waiting)
tr = 2 + len(ESTABLISHMENT)
ref.cell(tr, SUM_COL, 'TOTAL').font = Font(bold=True)
for off, idx in ((3, 3), (4, 4), (5, 6)):
    col = get_column_letter(SUM_COL + off)
    ref.cell(tr, SUM_COL + off, f'=SUM({col}2:{col}{tr-1})').font = Font(bold=True)
for c in range(1, 7):
    ref.column_dimensions[get_column_letter(c)].width = 58 if c == 6 else 18
for c in range(SUM_COL, SUM_COL + 6):
    ref.column_dimensions[get_column_letter(c)].width = 17
ref.freeze_panes = 'A2'

# ── Data validation ───────────────────────────────────────
MAXR = 1200
def add_list_dv(col_letter, values):
    dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"',
                        allow_blank=True, showDropDown=False)
    dv.error = 'Choose a value from the list.'; dv.errorTitle = 'Invalid entry'
    ws.add_data_validation(dv); dv.add(f'{col_letter}3:{col_letter}{MAXR}')

add_list_dv('A', LOCATIONS)
add_list_dv('H', GENDERS)
add_list_dv('I', ZONE_CODES)
add_list_dv('K', ZONE_CODES)
add_list_dv('L', STATUSES)
add_list_dv('M', DISCIPLINES)

# Bed codes: too many for an inline list, so point at the Reference range.
dv_bed = DataValidation(type='list', formula1=f'=Reference!$A$2:$A${len(inv)+1}',
                        allow_blank=True, showDropDown=False)
dv_bed.error = 'Use a position code from the Reference sheet.'
dv_bed.errorTitle = 'Unknown position code'
ws.add_data_validation(dv_bed); dv_bed.add(f'J3:J{MAXR}')

dv_age = DataValidation(type='whole', operator='between', formula1='0', formula2='120',
                        allow_blank=True)
dv_age.error = 'Age must be 0–120.'; dv_age.errorTitle = 'Invalid age'
ws.add_data_validation(dv_age); dv_age.add(f'G3:G{MAXR}')

# ── Conditional formatting ────────────────────────────────
ws.conditional_formatting.add(f'J3:J{MAXR}', FormulaRule(
    formula=['AND($J3<>"",ISNUMBER(SEARCH("crisis",$J3)))'],
    fill=PatternFill('solid', fgColor='FDE8E6'), font=Font(color='C0392B'), stopIfTrue=False))
ws.conditional_formatting.add(f'J3:J{MAXR}', FormulaRule(
    formula=['AND($J3<>"",ISNUMBER(SEARCH("waiting",$J3)))'],
    fill=PatternFill('solid', fgColor='E6F7EE'), font=Font(color='1A7C4A'), stopIfTrue=False))
for status, bg, fg in [('ongoingtreatment','E0F5F6','055257'), ('referred','E6F0FA','1A4C82'),
                       ('preadmit','FFF4E0','8A4B00'), ('admitted','E6F7EE','1A7C4A'),
                       ('discharge','F3EEFF','6B3EA0')]:
    ws.conditional_formatting.add(f'L3:L{MAXR}', FormulaRule(
        formula=[f'$L3="{status}"'], fill=PatternFill('solid', fgColor=bg),
        font=Font(color=fg), stopIfTrue=False))
# Admit recorded before triage — an impossible record.
ws.conditional_formatting.add(f'A3:U{MAXR}', FormulaRule(
    formula=['AND($B3<>"",$Q3<>"",$Q3<$B3)'],
    fill=PatternFill('solid', fgColor='FFE0E0'), stopIfTrue=False))

wb.save('ED_PAC_Register.xlsx')

# ── Report ────────────────────────────────────────────────
from collections import Counter
print(f'ED_PAC_Register.xlsx written: {n} records, {len(inv)} position codes')
print('records by location :', dict(Counter(r['Location'] for r in rows)))
print('records by status   :', dict(Counter(r['Status'] for r in rows)))
print('records by zone     :', dict(Counter(r['Zone Code'] for r in rows)))
print('waiting-area records:', sum(1 for r in rows if 'waiting' in r['Bed / Position Code']))
print('crisis-bed records  :', sum(1 for r in rows if 'crisis' in r['Bed / Position Code']))
print('triage window       :', min(r['Triage Date/Time'] for r in rows), '→',
      max(r['Triage Date/Time'] for r in rows))
print('arrivals per hour   :', dict(sorted(Counter(r['Triage Date/Time'].hour for r in rows).items())))
